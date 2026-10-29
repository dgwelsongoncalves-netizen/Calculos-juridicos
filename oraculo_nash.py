"""Oráculo: confere os NÚMEROS do motor (nash.py) E testa a Interface Gráfica (interface_v3.py) de ponta a ponta."""
import sys, io, json, math, random, tempfile, contextlib
from pathlib import Path
from unittest.mock import patch
import openpyxl, pandas as pd
import nash
nash.preload_heavy_libs(); nash.converter_para_pdf = lambda x: True
falhas = []
def falha(tag, msg): falhas.append((tag, msg))
def prox(a, b, tol=0.02): return abs(a-b) <= tol

# ---------- U: checagens unitárias dos motores, com reimplementação independente em Python puro ----------
tj = nash.carregar_tjmg()
cache = pd.read_csv(nash.ARQUIVO_CACHE_BCB, index_col=0, parse_dates=True)
bcb = {k: cache[[k]] for k in ['SELIC','IPCA','TAXA_LEGAL','IPCA_E','POUPANCA']}
T = pd.Timestamp
def meses(a, b):  # lista de inícios de mês de a até b (inclusive)
    return list(pd.date_range(T(a.year,a.month,1), T(b.year,b.month,1), freq='MS'))
def v(serie, col, m):
    try: x = float(serie.loc[m, col]); return None if math.isnan(x) else x
    except KeyError: return None

for ini, fim in [(T('2015-03-10'),T('2019-06-01')),(T('2019-06-01'),T('2024-08-30')),(T('2012-01-05'),T('2024-08-30'))]:
    f = nash.obter_fator_tjmg(tj, ini, fim)
    if f < 1: falha("U1","TJMG: fator %s->%s = %.6f (<1, correção encolhendo o valor)"%(ini.date(),fim.date(),f))
mid = T('2019-06-01'); a=T('2015-03-10'); c=T('2024-08-30')
if abs(nash.obter_fator_tjmg(tj,a,c) - nash.obter_fator_tjmg(tj,a,mid)*nash.obter_fator_tjmg(tj,mid,c)) > 1e-9:
    falha("U1b","TJMG: fator não é multiplicativo entre períodos")

for cm, jr, calc in [(T('2022-02-10'),T('2022-02-10'),T('2026-09-15')),(T('2021-05-03'),T('2023-01-20'),T('2026-09-15'))]:
    _, j = nash.calc_tjmg_juros(tj, cm, jr, calc)
    esp = max(0,(calc.year-jr.year)*12+(calc.month-jr.month))*0.01
    if not prox(j, esp, 1e-12): falha("U2","R1: juros %.4f != 1%%×meses (%.4f)"%(j,esp))

for cm, jr, calc in [(T('2024-09-10'),T('2024-09-10'),T('2026-09-15')),(T('2025-01-05'),T('2025-06-20'),T('2026-08-15'))]:
    fi, ju = nash.calc_leinova_pura(bcb, cm, jr, calc)
    esp_f = 1.0
    for m in meses(cm, calc):
        x = v(bcb['IPCA'],'IPCA',m); esp_f *= (1+x) if x is not None else 1.0
    esp_j = 0.0
    for m in meses(jr, calc):
        s, i = v(bcb['SELIC'],'SELIC',m), v(bcb['IPCA'],'IPCA',m)
        if s is not None and i is not None: esp_j += max(0.0, s-i)
    if not prox(fi, esp_f, 1e-9): falha("U3a","Lei 14.905: IPCA acumulado %.6f != produto esperado %.6f"%(fi,esp_f))
    if not prox(ju, esp_j, 1e-9): falha("U3b","Lei 14.905: juros %.6f != Σmax(0,SELIC-IPCA) esperado %.6f"%(ju,esp_j))

_, jp = nash.calc_fazenda_publica(bcb, T('2019-01-10'), T('2019-01-10'), T('2026-09-15'))
esp_p = sum(v(bcb['POUPANCA'],'POUPANCA',m) or 0.0 for m in meses(T('2019-01-10'), T('2021-11-01')))
if bcb['POUPANCA'].index.min() <= T('2019-01-01') and not prox(jp, esp_p, 1e-9): falha("U4","Fazenda: juros poupança %.6f != Σ esperado %.6f"%(jp,esp_p))
_, jp2 = nash.calc_fazenda_publica(bcb, T('2020-06-10'), T('2020-06-10'), T('2026-09-15'))
esp_p2 = sum(v(bcb['POUPANCA'],'POUPANCA',m) or 0.0 for m in meses(T('2020-06-10'), T('2021-11-01')))
if not prox(jp2, esp_p2, 1e-9): falha("U4","Fazenda: juros poupança %.6f != Σ esperado %.6f"%(jp2,esp_p2))

# ---------- E: cenários completos, lendo o laudo gerado ----------
def d(rng, a, b):
    ini = pd.Timestamp(a).value//10**9; fim = pd.Timestamp(b).value//10**9
    return pd.Timestamp(rng.randint(ini,fim), unit='s').strftime('%d/%m/%Y')
def escrever(path, cen):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Parametros"
    for k, val in cen["param"].items(): ws.append([k, val])
    w = wb.create_sheet("Danos"); w.append(["ID / Folha","Descrição","Data Desembolso","Valor Histórico","Regra","Valor Pedido Inicial","Data do Pedido","Data Juros"])
    for r in cen["danos"]: w.append(list(r) + [""] * (8 - len(r)))
    w = wb.create_sheet("Custas"); w.append(["ID / Folha","Descrição","Data Desembolso","Valor Histórico"])
    for r in cen["custas"]: w.append(r)
    w = wb.create_sheet("Deducoes"); w.append(["ID / Folha","Data bloqueio/deposito","Valor"])
    for r in cen["deducoes"]: w.append(r)
    wa = wb.create_sheet("Acordo_Params")
    for k_ac in ["Data do Inadimplemento", "Regra de Atualização Pós-Quebra", "Valor Original Confessado (Sem Desconto)"]:
        wa.append([k_ac, None])
    wb.save(path)

def gerar(rng, n):
    pag = rng.choice(["Sim","Não","Parcial","Não se aplica"])
    p = {"Processo":f"ORA-{n:03d}","Atuação":rng.choice(["RÉU","AUTOR"]),"Data do Trânsito":d(rng,'2024-01-01','2026-06-01') if rng.random()<.6 else "",
         "Data da Sentença":d(rng,'2023-01-01','2025-12-01'),"Termo Inicial Juros":rng.choice(["Desembolso","Citação","Evento"]),
         "Data da Citação":d(rng,'2021-01-01','2023-01-01'),"Data do Evento":d(rng,'2020-06-01','2022-01-01'),
         "Justiça Gratuita":rng.choice(["Não","Sim"]),"Juros sobre Custas":"Sim","Pagamento Voluntário 15d":pag,
         "Valor Pagamento Parcial (R$)":round(rng.uniform(100,2000),2) if pag=="Parcial" else "","Fazenda Pública":rng.choice(["Não","Não","Sim"]),
         "Honorários Sucumbência (%)":rng.choice([10,15,20]),"Honorários Fixos (R$)":"","Base Honorários":"CONDENAÇÃO",
         "Valor Causa Original":"","Data Propositura":d(rng,'2021-01-01','2022-01-01'),"Proporção Honorários (%)":100,"Proporção Custas (%)":100}
    uni = rng.random() < 0.5; r0 = rng.choice(["R1","R2","R3","R4","R5","R6","R7"])
    danos = [[f"F{i}",f"Dano {n}-{i}",d(rng,'2021-02-01','2025-06-01'),round(rng.uniform(300,9000),2),r0 if uni else rng.choice(["R1","R2","R3","R4","R5","R6","R7"]),"",""] for i in range(rng.randint(1,4))]
    custas = [[f"G{i}",f"Custa {n}-{i}",d(rng,'2021-02-01','2025-06-01'),round(rng.uniform(40,400),2)] for i in range(rng.randint(0,2))]
    ded = [[f"D{i}",d(rng,'2025-07-01','2026-06-01'),round(rng.uniform(300,4000),2)] for i in range(rng.choice([0,0,1,2,3]))]
    return {"param":p,"danos":danos,"custas":custas,"deducoes":ded}
def nums(row): return [c for c in row if isinstance(c,(int,float)) and not isinstance(c,bool)]
def txt(row): return " ".join(str(c) for c in row if c is not None)
def rodar(cen):
    with tempfile.TemporaryDirectory() as td:
        ent, sai = Path(td)/"e.xlsx", Path(td)/"s.xlsx"; escrever(ent, cen)
        with patch("bcb.sgs.get", side_effect=ConnectionError("offline")), contextlib.redirect_stdout(io.StringIO()):
            nash.executar_nash(str(ent), str(sai))
        for f in Path(td).iterdir():
            wb = openpyxl.load_workbook(f, data_only=True)
            if "Laudo de Liquidação" in wb.sheetnames: return [list(r) for r in wb["Laudo de Liquidação"].iter_rows(values_only=True)]
    return None
def checar(cen, n):
    tag = f"E[{n}]"; rows = rodar(cen); p = cen["param"]
    if rows is None: return falha(tag, "laudo não foi gerado")
    def val(rot):
        x = [nums(r)[-1] for r in rows if rot in txt(r) and nums(r)]; return x[0] if x else None
    sd, sc = val("Subtotal Danos Materiais"), val("Subtotal Custas")
    for dn in cen["danos"]:   # E1: nenhuma linha de dano some e o valor histórico aparece
        ok = any(dn[1] in txt(r) and any(prox(x, dn[3], 0.006) for x in nums(r)) for r in rows)
        if not ok: falha(tag, f"E1: dano '{dn[1]}' (R$ {dn[3]}) não aparece no laudo")
    fim = [r for r in rows if "TOTAL GERAL DEVIDO:" in txt(r) or "SALDO DEVEDOR FINAL DA EXECUÇÃO:" in txt(r)]
    if not fim: return falha(tag, "E6: linha do total final não encontrada")
    tot = nums(fim[-1])[-1]
    if not math.isfinite(tot) or tot < -0.01: falha(tag, f"E6: total inválido {tot}")
    fazenda = str(p["Fazenda Pública"]).upper()=="SIM"; jg = str(p["Justiça Gratuita"]).upper()=="SIM"
    multas = [nums(r)[0] for r in rows if "Multa Art. 523" in txt(r) and nums(r)]
    hons = [nums(r)[0] for r in rows if "Honorários Art. 523" in txt(r) and nums(r)]
    if fazenda and (any(x>0.005 for x in multas) or any(x>0.005 for x in hons)): falha(tag,"E2: Fazenda Pública com multa/honorários do art. 523")
    if jg and any(x>0.005 for x in hons): falha(tag,"E2: Justiça gratuita com honorários do art. 523")
    if multas and hons and not jg and not prox(multas[0], hons[0], 0.02): falha(tag,f"E2: multa (R$ {multas[0]:.2f}) != honorários 523 (R$ {hons[0]:.2f}); ambos deveriam ser 10% da mesma base")
    if not cen["deducoes"]:   # E8/E5/E11: aritmética do resumo e base de honorários (Art. 85 § 2º CPC)
        st = val("SUBTOTAL (Principal")
        if None not in (sd, sc, st) and not prox(st, sd+sc, 0.05): falha(tag, f"E8: SUBTOTAL {st:.2f} != danos {sd:.2f} + custas {sc:.2f}")
        h = val("Honorários Acumulados") or 0.0
        esp_h = 0.0 if jg else (sd or 0.0) * (float(p["Honorários Sucumbência (%)"]) / 100.0)
        if sd is not None and not prox(h, esp_h, 0.05): falha(tag, f"E11: Honorários {h:.2f} != {p['Honorários Sucumbência (%)']}% sobre Danos ({esp_h:.2f}) - possível incidência indevida sobre custas")
        if str(p["Pagamento Voluntário 15d"]).upper()=="PARCIAL" and not fazenda and multas:
            rem = val("SALDO REMANESCENTE")
            if rem is not None and not prox(multas[0], 0.10*rem, 0.02): falha(tag, f"E9: multa 523 (R$ {multas[0]:.2f}) != 10% do saldo remanescente (R$ {0.10*rem:.2f})")
        if str(p["Pagamento Voluntário 15d"]).upper()!="PARCIAL":
            m = multas[0] if multas else 0.0; h5 = hons[0] if hons else 0.0
            if st is not None and not prox(tot, st+h+m+h5, 0.05): falha(tag, f"E5: total {tot:.2f} != soma das parcelas {st+h+m+h5:.2f}")
    else:                      # E3/E4/E12: conta gráfica
        linha_ini_cg = [nums(r) for r in rows if "Atualização até 1º Bloqueio" in txt(r) and len(nums(r)) >= 3]
        if linha_ini_cg and None not in (sd, sc):
            saldo_pre_bloq = linha_ini_cg[0][-1]
            if saldo_pre_bloq > (sd + sc) * 1.01 + 0.05:
                falha(tag, f"E12: Saldo no 1º bloqueio (R$ {saldo_pre_bloq:.2f}) maior que Danos+Custas até hoje (R$ {sd+sc:.2f}) - duplicação na Conta Gráfica!")
        dep = [i for i,r in enumerate(rows) if "Bloqueio/Depósito" in txt(r) and "ID:" in txt(r)]
        if len(dep) != len(cen["deducoes"]): falha(tag, f"E3: {len(dep)} bloqueios no laudo, {len(cen['deducoes'])} informados")
        decl = sorted(cen["deducoes"], key=lambda r: pd.to_datetime(r[1], dayfirst=True))
        for k, i in enumerate(dep):
            if k >= len(decl): break
            now, prev = nums(rows[i]), nums(rows[i-1])
            if len(now) < 4 or len(prev) < 2: continue
            P0, J0 = prev[0], prev[1]; vd = decl[k][2]; P1, J1, pago = now[0], now[1], now[2]
            if not prox(pago, vd, 0.01): falha(tag, f"E3: valor do bloqueio no laudo R$ {pago:.2f} != informado R$ {vd:.2f}")
            eJ = max(0.0, J0-vd); eP = max(0.0, P0-max(0.0, vd-J0))
            if not (prox(J1, eJ, 0.05) and prox(P1, eP, 0.05)): falha(tag, f"E3: imputação do bloqueio fora do art. 354 CC (juros primeiro). Esperado P={eP:.2f} J={eJ:.2f}; laudo P={P1:.2f} J={J1:.2f}")
        uniforme = len({x[4] for x in cen["danos"]}) == 1 and not cen["custas"]
        for i, r in enumerate(rows):
            if uniforme and ("Atualização do Saldo" in txt(r) or "Atualização Final" in txt(r)) and i>0:
                a, b = nums(rows[i-1]), nums(r)
                if len(a)>=2 and len(b)>=2 and a[0]>1 and a[1]>1 and b[1] < a[1]*(b[0]/a[0])*(1-1e-7): falha(tag, f"E4: juros acumulados não acompanham a correção do principal (J {a[1]:.2f}->{b[1]:.2f}, P {a[0]:.2f}->{b[0]:.2f})")
        if len(cen["danos"]) >= 2:   # E10: ordem das linhas não pode mudar o total
            c2 = dict(cen); c2["danos"] = list(reversed(cen["danos"])); rows2 = rodar(c2)
            f2 = [r for r in rows2 if "SALDO DEVEDOR FINAL DA EXECUÇÃO:" in txt(r)] if rows2 else []
            if f2 and not prox(tot, nums(f2[-1])[-1], 0.02): falha(tag, f"E10: total muda com a ordem das linhas de Danos ({tot:.2f} x {nums(f2[-1])[-1]:.2f})")

def checar_entradas_invalidas():
    """V: dado ilegível tem que gerar ERRO EXPLÍCITO. Nunca pode ser descartado nem virar zero em silêncio."""
    base = lambda: {"param": gerar(random.Random(5), 0)["param"], "danos": [["F1","Dano V","10/10/2024",5000.0,"R6","",""]], "custas": [], "deducoes": [["D1","10/01/2026",1000.0]]}
    casos = [("bloqueio com data inexistente (31/02)", lambda c: c["deducoes"].__setitem__(0, ["D1","31/02/2026",1000.0])),
             ("bloqueio com data em texto", lambda c: c["deducoes"].__setitem__(0, ["D1","ontem",1000.0])),
             ("bloqueio sem valor", lambda c: c["deducoes"].__setitem__(0, ["D1","10/01/2026",None])),
             ("dano com data inexistente (31/02)", lambda c: c["danos"].append(["F2","Dano X","31/02/2024",3000.0,"R6","",""])),
             ("dano com valor ilegível (abc)", lambda c: c["danos"].append(["F2","Dano X","10/11/2024","abc","R6","",""])),
             ("dano com valor e sem data", lambda c: c["danos"].append(["F2","Dano X","",3000.0,"R6","",""])),
             ("dano com 'Data Juros' inexistente (31/02)", lambda c: c["danos"].append(["F2","Dano X","10/11/2024",3000.0,"R6","","","31/02/2025"])),
             ("dano com 'Data do Pedido' inexistente (31/02)", lambda c: c["danos"].append(["F2","Dano X","10/11/2024",3000.0,"R6",4000.0,"31/02/2022"])),
             ("data como número do Excel (44966)", lambda c: c["danos"].__setitem__(0, ["F1","Dano V",44966,5000.0,"R6","",""])),
             ("custa com valor e sem data", lambda c: c["custas"].append(["G1","Custa X","",100.0]))]
    for nome, alterar in casos:
        c = base(); alterar(c)
        try: rodar(c); falha("V", f"entrada inválida aceita sem erro: {nome}")
        except Exception: pass

def checar_interface_gui():
    """G: Testa a interface_v3.py automaticamente (Modos Guiados, Inserção Rápida, Colar do Excel, Auto-Save e Geração de Cálculo)."""
    if not Path("interface_v3.py").exists():
        return
    try:
        import tkinter as tk
        import interface_v3
    except Exception as e:
        return falha("G0", f"Erro ao importar interface_v3.py: {e}")

    try:
        root = tk.Tk()
        root.withdraw()  # Roda em modo invisível para não travar a tela
    except Exception:
        return  # Ambiente sem servidor gráfico X11/Wayland

    with tempfile.TemporaryDirectory() as td:
        rascunho_temp = Path(td) / "rascunho_teste.xlsx"
        # Isola o rascunho de teste para não sobrescrever o rascunho real do usuário e desativa rede do updater
        with patch.object(interface_v3, "ARQUIVO_RASCUNHO", rascunho_temp), \
             patch.object(interface_v3.NashDataEntryGUI, "verificar_atualizacao_background", lambda self, silencioso=True: None), \
             patch("tkinter.messagebox.showinfo"), \
             patch("tkinter.messagebox.showwarning"):

            app = interface_v3.NashDataEntryGUI(root)

            # G1: Teste do Modo 2 (Execução de Título Extrajudicial)
            app.vars_param["Data da Sentença"].set("10/05/2024")
            app.vars_param["Data do Trânsito"].set("10/06/2024")
            app.vars_param["Pagamento Voluntário 15d"].set("Não")
            app.modo_calculo.set("EXECUCAO_EXTRA")
            app.aplicar_modo_calculo()

            if app.vars_param["Data da Sentença"].get() != "" or app.vars_param["Data do Trânsito"].get() != "":
                falha("G1", "Modo Execução Extrajudicial não limpou Data da Sentença / Trânsito")
            if app.vars_param["Pagamento Voluntário 15d"].get() != "Não se aplica":
                falha("G1", "Modo Execução Extrajudicial não bloqueou Multa Art. 523")
            if app.vars_param["Termo Inicial Juros"].get() != "Desembolso":
                falha("G1", "Modo Execução Extrajudicial não ajustou Termo Inicial Juros para Desembolso")

            # G2: Teste de Inserção Rápida Inline (Danos, Custas e Bloqueios) — Reproduzindo processo 5047942-62.2022.8.13.0024
            app.vars_param["Processo"].set("5047942-62.2022.8.13.0024")
            app.vars_param["Atuação"].set("AUTOR")
            app.vars_param["Data Propositura"].set("26/05/2022")
            app.vars_param["Honorários Sucumbência (%)"].set("10")

            app.in_dano_id.insert(0, "Evento 1")
            app.in_dano_desc.insert(0, "ACORDO EXTRAJUDICIAL")
            app.in_dano_data.insert(0, "14/12/2021")
            app.in_dano_val.insert(0, "1859,40")
            app.in_dano_regra.set("R1 - TJMG + Juros de 1% a.m.")
            app.adicionar_inline_dano()

            if len(app.tree_danos.get_children()) != 1:
                falha("G2", "Inserção rápida de Dano não adicionou linha na tabela")

            custas_teste = [
                ("Ev 15", "CUSTAS CITACAO", "14/04/2022", "36,60"),
                ("Ev 20", "CUSTAS CITACAO", "05/05/2022", "35,96"),
                ("Ev 68", "4x OFICIOS", "09/10/2023", "69,00"),
                ("Ev 69a", "CUSTAS ALVARA", "15/09/2023", "60,44"),
                ("Ev 69b", "CUSTAS PENHORA", "15/09/2023", "76,61"),
                ("Ev 105", "CUSTAS SISBAJUD", "24/10/2025", "13,27"),
                ("Ev 123", "ALVARA POS BLOQUEIO", "16/07/2026", "83,38"),
            ]
            for cid, cdesc, cdt, cval in custas_teste:
                app.in_custa_id.insert(0, cid)
                app.in_custa_desc.insert(0, cdesc)
                app.in_custa_data.insert(0, cdt)
                app.in_custa_val.insert(0, cval)
                app.adicionar_inline_custa()

            if len(app.tree_custas.get_children()) != 7:
                falha("G2", f"Inserção rápida de Custas: esperado 7, obteve {len(app.tree_custas.get_children())}")

            # G3: Teste de Colagem do Excel (Ctrl+V) na aba Deduções
            texto_tsv = (
                "Evento 128, SISBAJUD1\t07/05/2026\t1492,86\n"
                "Evento 128, SISBAJUD2\t30/04/2026\t1557,80\n"
                "Evento 128, SISBAJUD3\t24/04/2026\t0,01\n"
                "Evento 128, SISBAJUD4\t22/04/2026\t5,04\n"
            )
            with patch.object(root, "clipboard_get", return_value=texto_tsv):
                app.colar_do_excel("deducoes")

            if len(app.tree_deducoes.get_children()) != 4:
                falha("G3", f"Colagem do Excel em Deduções: esperado 4 linhas, obteve {len(app.tree_deducoes.get_children())}")

            # G4: Conferência Rápida ("Conta de Padaria") e Auto-Save
            app.atualizar_resumo_rapido()
            texto_resumo = app.lbl_resumo_valores.cget("text")
            if "1,859.40" not in texto_resumo and "1859.40" not in texto_resumo:
                falha("G4", f"Resumo rápido não exibiu soma correta dos danos: {texto_resumo}")
            if not rascunho_temp.exists():
                falha("G4", "Auto-Save não gerou o arquivo de rascunho automático")

            # G5: Exportação do Template pela Interface -> Reimportação -> Execução no Motor Matemático
            arq_template_gui = Path(td) / "template_gerado_pela_gui.xlsx"
            arq_laudo_gui = Path(td) / "laudo_gerado_pela_gui.xlsx"
            app._gravar_workbook_template(arq_template_gui)

            # Limpa uma tabela e reimporta o arquivo salvo para garantir fidelidade 100% de Ida e Volta
            app.limpar_tree("danos")
            app._carregar_workbook_na_tela(arq_template_gui, silencioso=True)
            if len(app.tree_danos.get_children()) != 1 or len(app.tree_custas.get_children()) != 7 or len(app.tree_deducoes.get_children()) != 4:
                falha("G5", "Perda de linhas ao exportar e reimportar template pela interface")

            with patch("bcb.sgs.get", side_effect=ConnectionError("offline")), contextlib.redirect_stdout(io.StringIO()):
                nash.executar_nash(str(arq_template_gui), str(arq_laudo_gui))

            wb_out = openpyxl.load_workbook(arq_laudo_gui, data_only=True)
            rows_out = [list(r) for r in wb_out["Laudo de Liquidação"].iter_rows(values_only=True)]
            fim_out = [r for r in rows_out if "SALDO DEVEDOR FINAL DA EXECUÇÃO:" in txt(r)]
            if not fim_out:
                falha("G5", "Laudo gerado a partir do template da interface não tem linha de saldo final")
            else:
                saldo_final_gui = nums(fim_out[-1])[-1]
                # O saldo real do caso 5047942-62.2022.8.13.0024 deve ficar na faixa de R$ 1.200 a R$ 1.450, jamais R$ 5.400!
                if not (1150.0 <= saldo_final_gui <= 1500.0):
                    falha("G5", f"Saldo final do caso 5047942 via GUI fora do esperado: R$ {saldo_final_gui:.2f}")

        root.destroy()

rng = random.Random(7)
for n in range(1, 81): checar(gerar(rng, n), n)
checar_entradas_invalidas()
checar_interface_gui()
print(json.dumps({"n_falhas": len(falhas), "tags": sorted({(m.split(':')[0] if m.startswith('E') else t) for t,m in falhas}), "amostra": [f"{t} {m}"[:150] for t,m in falhas[:3]]}, ensure_ascii=False))