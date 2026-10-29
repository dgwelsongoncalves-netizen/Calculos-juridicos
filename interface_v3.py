import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from datetime import datetime
import subprocess
import sys
import os
import re
import json
import tempfile
import urllib.request
from pathlib import Path
import threading
import nash

GITHUB_REPO = "dgwelsongoncalves-netizen/Calculos-juridicos"
ARQUIVO_RASCUNHO = Path.home() / ".nash_rascunho_auto.xlsx"


class NashDataEntryGUI:
    def __init__(self, root):
        self.root = root
        versao_atual = getattr(nash, "__version__", "4.1.3")
        self.root.title(f"Nash System v{versao_atual} — Liquidação Judicial Assistida")
        self.root.geometry("1040x740")
        self.root.minsize(960, 680)

        # Opções de Regras Globais
        self.opcoes_regras = [
            "R1 - TJMG + Juros de 1% a.m.",
            "R2 - Taxa Selic (critério único)",
            "R3 - TJMG + Juros 1% a.m. até 08/24; após, Selic",
            "R4 - TJMG + Juros 1% a.m. até 08/24; após, Lei 14.905/24",
            "R5 - Selic até 08/24; após, Lei 14.905/24 (Tema 1.368 STJ)",
            "R6 - Lei 14.905/24: IPCA + Taxa Legal",
            "R7 - TJMG CM + Taxa Legal Retroativa"
        ]

        self.modo_calculo = tk.StringVar(value="SENTENCA")
        self.widgets_param = {}
        self.vars_param = {}
        self.vars_acordo = {}

        self._montar_layout_principal(versao_atual)
        self.aplicar_modo_calculo()

        # Restaura rascunho automático se existir
        self.root.after(200, self._tentar_restaurar_rascunho)
        # Verifica atualização silenciosamente em background (Padrão Windows)
        self.root.after(1500, lambda: self.verificar_atualizacao_background(silencioso=True))
        # Proteção ao fechar a janela no "X"
        self.root.protocol("WM_DELETE_WINDOW", self.ao_fechar_janela)

    def _montar_layout_principal(self, versao_atual):
        # 1. Cabeçalho Superior Escuro
        header = tk.Frame(self.root, bg="#1F3864", padx=14, pady=8)
        header.pack(fill="x", side="top")

        tk.Label(
            header,
            text=f"⚖️  NASH SYSTEM v{versao_atual} — Liquidação Judicial Assistida",
            bg="#1F3864", fg="white", font=("Arial", 12, "bold")
        ).pack(side="left")

        btn_update = tk.Button(
            header,
            text="🔄 Verificar Atualização",
            bg="#334155", fg="white", activebackground="#475569", activeforeground="white",
            relief="flat", font=("Arial", 8, "bold"), padx=8, pady=2,
            command=lambda: self.verificar_atualizacao_background(silencioso=False)
        )
        btn_update.pack(side="right", padx=(10, 0))

        self.lbl_autosave = tk.Label(
            header,
            text="● Auto-Save Ativo",
            bg="#1F3864", fg="#86EFAC", font=("Arial", 9, "bold")
        )
        self.lbl_autosave.pack(side="right", padx=10)

        # 2. Container Principal
        main_frame = ttk.Frame(self.root, padding=10)
        main_frame.pack(fill="both", expand=True)

        # 3. Rodapé Fixo (Conferência Rápida + Botões de Ação)
        bottom_container = ttk.Frame(main_frame)
        bottom_container.pack(fill="x", side="bottom", pady=(6, 0))

        # Caixa "Conta de Padaria" (Prova dos Nove em tempo real)
        self.frame_resumo = tk.Frame(bottom_container, bg="#EFF6FF", highlightbackground="#2563EB", highlightthickness=1, padx=10, pady=6)
        self.frame_resumo.pack(fill="x", pady=(0, 8))

        tk.Label(
            self.frame_resumo,
            text="🔍 CONFERÊNCIA RÁPIDA EM TEMPO REAL ('CONTA DE PADARIA' ANTES DE EXPORTAR):",
            bg="#EFF6FF", fg="#1D4ED8", font=("Arial", 8, "bold")
        ).pack(anchor="w")

        self.lbl_resumo_valores = tk.Label(
            self.frame_resumo,
            text="Dívida Histórica: R$ 0,00   -->   (+) Custas Lançadas: R$ 0,00   -->   (-) Bloqueios/Depósitos: -R$ 0,00",
            bg="#EFF6FF", fg="#0F172A", font=("Arial", 9, "bold")
        )
        self.lbl_resumo_valores.pack(anchor="w", pady=(2, 0))

        # Barra de Botões Principais
        action_frame = ttk.Frame(bottom_container)
        action_frame.pack(fill="x")

        ttk.Button(action_frame, text="📂 Importar Template Excel", command=self.importar_planilha).pack(side="left", padx=(0, 6))
        ttk.Button(action_frame, text="⚡ Calcular Direto de Planilha", command=self.calcular_direto_planilha).pack(side="left", padx=6)
        ttk.Button(action_frame, text="🧹 Limpar / Novo Cálculo", command=self.novo_calculo_limpo).pack(side="left", padx=6)

        self.btn_processar = tk.Button(
            action_frame,
            text="🚀 Gerar Template + Laudo PDF",
            bg="#16A34A", fg="white", activebackground="#15803D", activeforeground="white",
            font=("Arial", 9, "bold"), relief="flat", padx=14, pady=5,
            command=self.salvar_e_gerar
        )
        self.btn_processar.pack(side="right", padx=(6, 0))

        btn_so_template = tk.Button(
            action_frame,
            text="💾 Só Salvar Template (.xlsx)",
            bg="#0284C7", fg="white", activebackground="#0369A1", activeforeground="white",
            font=("Arial", 9, "bold"), relief="flat", padx=12, pady=5,
            command=self.salvar_apenas_template
        )
        btn_so_template.pack(side="right", padx=6)

        # 4. Notebook (Abas)
        self.notebook = ttk.Notebook(main_frame)
        self.notebook.pack(fill="both", expand=True)

        self.tab_parametros = ttk.Frame(self.notebook, padding=10)
        self.tab_danos = ttk.Frame(self.notebook, padding=10)
        self.tab_custas = ttk.Frame(self.notebook, padding=10)
        self.tab_deducoes = ttk.Frame(self.notebook, padding=10)
        self.tab_acordos = ttk.Frame(self.notebook, padding=10)
        self.tab_dashboard = ttk.Frame(self.notebook, padding=10)

        self.notebook.add(self.tab_parametros, text=" 1. Tipo de Ação e Parâmetros ")
        self.notebook.add(self.tab_danos, text=" 2. Valores / Danos (0) ")
        self.notebook.add(self.tab_custas, text=" 3. Custas (0) ")
        self.notebook.add(self.tab_deducoes, text=" 4. Bloqueios / Depósitos (0) ")
        self.notebook.add(self.tab_acordos, text=" 5. Quebra de Acordo ")
        self.notebook.add(self.tab_dashboard, text=" 6. Visão Estratégica ")

        self.montar_aba_parametros()
        self.montar_aba_danos()
        self.montar_aba_custas()
        self.montar_aba_deducoes()
        self.montar_aba_acordos()
        self.montar_aba_dashboard()

    # ================= ABA 1: PARÂMETROS GUIADOS =================
    def montar_aba_parametros(self):
        # PASSO 1: Seletor de Natureza do Cálculo
        frame_modo = tk.LabelFrame(
            self.tab_parametros,
            text=" PASSO 1 — Qual é a natureza deste cálculo? (Configura e bloqueia campos desnecessários automaticamente) ",
            font=("Arial", 9, "bold"), fg="#1F3864", padx=10, pady=8
        )
        frame_modo.pack(fill="x", pady=(0, 10))

        modos = [
            ("SENTENCA", "1. Cumprimento de Sentença (Condenação Judicial)", "Calcula Danos, Custas, Honorários Sucumbenciais e Multa Art. 523."),
            ("EXECUCAO_EXTRA", "2. Execução de Título Extrajudicial (Acordo / Contrato)", "Juros desde o vencimento + Hon. Art. 827 (Bloqueia Sentença/Trânsito/523)."),
            ("QUEBRA_ACORDO", "3. Quebra de Acordo Homologado (Cláusula Penal)", "Reverte desconto, abate parcelas pagas na Aba 5 e aplica multa contratual.")
        ]

        for idx, (val_modo, titulo, subtitulo) in enumerate(modos):
            f_card = tk.Frame(frame_modo, bd=1, relief="groove", padx=8, pady=4)
            f_card.grid(row=0, column=idx, sticky="nsew", padx=5)
            frame_modo.columnconfigure(idx, weight=1)

            rb = ttk.Radiobutton(
                f_card, text=titulo, variable=self.modo_calculo, value=val_modo,
                command=self.aplicar_modo_calculo
            )
            rb.pack(anchor="w")
            tk.Label(f_card, text=subtitulo, fg="#475569", font=("Arial", 8), justify="left", wraplength=280).pack(anchor="w", pady=(2, 0))

        # Container das 2 Colunas (Passo 2 e Passo 3)
        frame_cols = ttk.Frame(self.tab_parametros)
        frame_cols.pack(fill="both", expand=True)
        frame_cols.columnconfigure(0, weight=1)
        frame_cols.columnconfigure(1, weight=1)

        # PASSO 2 (Coluna Esquerda): Identificação e Datas
        frame_esq = tk.LabelFrame(
            frame_cols, text=" PASSO 2 — Identificação e Datas do Processo ",
            font=("Arial", 9, "bold"), fg="#1F3864", padx=10, pady=8
        )
        frame_esq.grid(row=0, column=0, sticky="nsew", padx=(0, 6))

        # PASSO 3 (Coluna Direita): Honorários, Custas e Art. 523
        frame_dir = tk.LabelFrame(
            frame_cols, text=" PASSO 3 — Honorários, Custas e Cumprimento (Art. 523) ",
            font=("Arial", 9, "bold"), fg="#1F3864", padx=10, pady=8
        )
        frame_dir.grid(row=0, column=1, sticky="nsew", padx=(6, 0))

        campos_esquerda = [
            ("Processo", "Entry", "", "Número CNJ do Processo:"),
            ("Atuação", "Combo", ["AUTOR", "RÉU"], "Nossa Atuação no Processo:"),
            ("Termo Inicial Juros", "Combo", ["Desembolso", "Citação", "Evento", "Propositura", "Trânsito em julgado"], "Juros de Mora contam a partir de:"),
            ("Data Propositura", "Entry", "", "Data da Propositura / Distribuição:"),
            ("Data da Citação", "Entry", "", "Data da Citação (se juros da citação):"),
            ("Data do Evento", "Entry", "", "Data do Evento Danoso (se Súm. 54):"),
            ("Data da Sentença", "Entry", "", "Data da Sentença:"),
            ("Data do Trânsito", "Entry", "", "Data do Trânsito em Julgado:"),
            ("Fazenda Pública", "Combo", ["Não", "Sim"], "Contra a Fazenda Pública (EC 113)?")
        ]

        campos_direita = [
            ("Base Honorários", "Combo", ["CONDENAÇÃO", "VALOR DA CAUSA", "PROVEITO ECONÔMICO"], "Base de Cálculo dos Honorários:"),
            ("Honorários Sucumbência (%)", "Entry", "10", "Percentual de Honorários (%):"),
            ("Honorários Fixos (R$)", "Entry", "", "Ou Honorários Fixos em R$ (opcional):"),
            ("Valor Causa Original", "Entry", "", "Valor da Causa Original (R$):"),
            ("Justiça Gratuita", "Combo", ["Não", "Sim"], "Parte Adversa tem Justiça Gratuita?"),
            ("Juros sobre Custas", "Combo", ["Não (Sem Juros)", "Sim (Após Trânsito)"], "Cobrar Juros de Mora sobre Custas?"),
            ("Pagamento Voluntário 15d", "Combo", ["Não se aplica", "Sim", "Não", "Parcial"], "Pagou em 15 dias (Multa Art. 523)?"),
            ("Valor Pagamento Parcial (R$)", "Entry", "", "Valor Pago Parcialmente no 523 (R$):"),
            ("Proporção Honorários (%)", "Entry", "100%", "Proporção de Rateio Honorários (%):"),
            ("Proporção Custas (%)", "Entry", "100%", "Proporção de Rateio Custas (%):")
        ]

        self._criar_campos_grade(frame_esq, campos_esquerda)
        self._criar_campos_grade(frame_dir, campos_direita)

        # Dica explicativa sobre "Desembolso"
        lbl_dica = tk.Label(
            frame_esq,
            text="💡 Dica: Em 'Desembolso', o Nash usa a data de cada item lançado na Aba 2 (ideal para Notas Fiscais ou Vencimento de Acordo).",
            fg="#1D4ED8", font=("Arial", 8, "italic"), wraplength=430, justify="left"
        )
        lbl_dica.grid(row=len(campos_esquerda), column=0, columnspan=2, sticky="w", pady=(8, 0))

    def _criar_campos_grade(self, parent, lista_campos):
        parent.columnconfigure(1, weight=1)
        for r_idx, (chave_interna, tipo, default, rotulo_amigavel) in enumerate(lista_campos):
            var = tk.StringVar()
            var.trace_add("write", lambda *args: self._ao_modificar_dados())
            self.vars_param[chave_interna] = var

            lbl = ttk.Label(parent, text=rotulo_amigavel)
            lbl.grid(row=r_idx, column=0, sticky="w", padx=4, pady=3)

            if tipo == "Entry":
                widget = ttk.Entry(parent, textvariable=var, width=26)
                if isinstance(default, str) and default:
                    var.set(default)
            else:
                widget = ttk.Combobox(parent, textvariable=var, values=default, state="readonly", width=24)
                var.set(default[0])

            widget.grid(row=r_idx, column=1, sticky="ew", padx=4, pady=3)
            self.widgets_param[chave_interna] = widget

    def aplicar_modo_calculo(self):
        modo = self.modo_calculo.get()
        if modo == "EXECUCAO_EXTRA":
            # Configura automaticamente para Execução de Título Extrajudicial
            self.vars_param["Termo Inicial Juros"].set("Desembolso")
            self.vars_param["Data da Sentença"].set("")
            self.vars_param["Data do Trânsito"].set("")
            self.vars_param["Pagamento Voluntário 15d"].set("Não se aplica")
            self.vars_param["Valor Pagamento Parcial (R$)"].set("")
            self.vars_param["Juros sobre Custas"].set("Não (Sem Juros)")

            self.widgets_param["Data da Sentença"].config(state="disabled")
            self.widgets_param["Data do Trânsito"].config(state="disabled")
            self.widgets_param["Pagamento Voluntário 15d"].config(state="disabled")
            self.widgets_param["Valor Pagamento Parcial (R$)"].config(state="disabled")
        else:
            self.widgets_param["Data da Sentença"].config(state="normal")
            self.widgets_param["Data do Trânsito"].config(state="normal")
            self.widgets_param["Pagamento Voluntário 15d"].config(state="readonly")
            self.widgets_param["Valor Pagamento Parcial (R$)"].config(state="normal")
            if modo == "QUEBRA_ACORDO":
                self.notebook.select(self.tab_acordos)

        self.atualizar_resumo_rapido()

    # ================= ABA 2: DANOS (COM INSERÇÃO RÁPIDA) =================
    def montar_aba_danos(self):
        f_inline = tk.LabelFrame(
            self.tab_danos,
            text=" ⚡ Inserção Rápida de Dano / Título Principal (Preencha e aperte [Enter] — sem precisar abrir janelas!) ",
            font=("Arial", 9, "bold"), fg="#1D4ED8", padx=8, pady=6
        )
        f_inline.pack(fill="x", pady=(0, 8))

        self.in_dano_id = ttk.Entry(f_inline, width=16)
        self.in_dano_desc = ttk.Entry(f_inline, width=28)
        self.in_dano_data = ttk.Entry(f_inline, width=12)
        self.in_dano_val = ttk.Entry(f_inline, width=13)
        self.in_dano_regra = ttk.Combobox(f_inline, values=self.opcoes_regras, state="readonly", width=26)
        self.in_dano_regra.current(0)

        campos_in = [
            ("ID / Evento:", self.in_dano_id),
            ("Descrição (Ex: Nota Fiscal / Acordo):", self.in_dano_desc),
            ("Data (DD/MM/AAAA):", self.in_dano_data),
            ("Valor (R$):", self.in_dano_val),
            ("Regra de Correção:", self.in_dano_regra)
        ]
        for idx, (rot, w) in enumerate(campos_in):
            tk.Label(f_inline, text=rot, font=("Arial", 8, "bold")).grid(row=0, column=idx, sticky="w", padx=4)
            w.grid(row=1, column=idx, sticky="ew", padx=4, pady=2)
            w.bind("<Return>", lambda e: self.adicionar_inline_dano())

        btn_inc = tk.Button(
            f_inline, text="➕ Incluir [Enter]", bg="#16A34A", fg="white",
            font=("Arial", 8, "bold"), relief="flat", padx=10, pady=3,
            command=self.adicionar_inline_dano
        )
        btn_inc.grid(row=1, column=len(campos_in), padx=6, pady=2)

        frame_botoes = ttk.Frame(self.tab_danos)
        frame_botoes.pack(fill="x", pady=(0, 6))

        ttk.Button(frame_botoes, text="➕ Completo (Com Pedido/Data Juros)", command=lambda: self.popup_generico('danos')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="✏️ Editar Selecionado", command=lambda: self.editar_selecionado('danos')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="🗑️ Remover", command=lambda: self.remover_selecionado('danos')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="📋 Colar do Excel", command=lambda: self.colar_do_excel('danos')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="🧹 Limpar Aba", command=lambda: self.limpar_tree('danos')).pack(side="left", padx=3)

        colunas = ('id', 'descricao', 'data_desembolso', 'valor', 'regra', 'data_juros', 'valor_pedido', 'data_pedido')
        self.tree_danos = ttk.Treeview(self.tab_danos, columns=colunas, show='headings', height=11)
        titulos = ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico', 'Regra', 'Data Juros', 'Valor Ped. Inicial', 'Data do Pedido']
        for col, titulo in zip(colunas, titulos):
            self.tree_danos.heading(col, text=titulo)
            self.tree_danos.column(col, width=105, anchor="center")
        self.tree_danos.column('descricao', width=230, anchor="w")
        self.tree_danos.bind("<Double-1>", lambda e: self.editar_selecionado('danos'))

        scroll_y = ttk.Scrollbar(self.tab_danos, orient="vertical", command=self.tree_danos.yview)
        self.tree_danos.configure(yscrollcommand=scroll_y.set)
        scroll_y.pack(side="right", fill="y")
        self.tree_danos.pack(fill="both", expand=True)

    def adicionar_inline_dano(self):
        desc = self.in_dano_desc.get().strip()
        data = self.in_dano_data.get().strip()
        val = self.in_dano_val.get().strip()
        if not desc or not data or not val:
            messagebox.showwarning("Campos Obrigatórios", "Preencha Descrição, Data e Valor para incluir o item.")
            return
        regra = self.in_dano_regra.get().split(" - ")[0].strip()
        self.tree_danos.insert('', 'end', values=(
            self.in_dano_id.get().strip(), desc, data, val, regra, "", "", ""
        ))
        self.in_dano_id.delete(0, 'end')
        self.in_dano_desc.delete(0, 'end')
        self.in_dano_data.delete(0, 'end')
        self.in_dano_val.delete(0, 'end')
        self.in_dano_id.focus_set()
        self._ao_modificar_dados()

    # ================= ABA 3: CUSTAS (COM INSERÇÃO RÁPIDA) =================
    def montar_aba_custas(self):
        f_inline = tk.LabelFrame(
            self.tab_custas,
            text=" ⚡ Inserção Rápida de Custa Processual (Digite e pressione [Enter]) ",
            font=("Arial", 9, "bold"), fg="#1D4ED8", padx=8, pady=6
        )
        f_inline.pack(fill="x", pady=(0, 8))

        self.in_custa_id = ttk.Entry(f_inline, width=22)
        self.in_custa_desc = ttk.Entry(f_inline, width=36)
        self.in_custa_data = ttk.Entry(f_inline, width=15)
        self.in_custa_val = ttk.Entry(f_inline, width=15)

        campos_c = [
            ("ID / Evento PJe:", self.in_custa_id),
            ("Descrição da Despesa (Ex: Custas Iniciais / Alvará):", self.in_custa_desc),
            ("Data Pagamento:", self.in_custa_data),
            ("Valor Pago (R$):", self.in_custa_val)
        ]
        for idx, (rot, w) in enumerate(campos_c):
            tk.Label(f_inline, text=rot, font=("Arial", 8, "bold")).grid(row=0, column=idx, sticky="w", padx=4)
            w.grid(row=1, column=idx, sticky="ew", padx=4, pady=2)
            w.bind("<Return>", lambda e: self.adicionar_inline_custa())

        btn_inc = tk.Button(
            f_inline, text="➕ Incluir [Enter]", bg="#16A34A", fg="white",
            font=("Arial", 8, "bold"), relief="flat", padx=10, pady=3,
            command=self.adicionar_inline_custa
        )
        btn_inc.grid(row=1, column=len(campos_c), padx=6, pady=2)

        frame_botoes = ttk.Frame(self.tab_custas)
        frame_botoes.pack(fill="x", pady=(0, 6))

        ttk.Button(frame_botoes, text="✏️ Editar Selecionada", command=lambda: self.editar_selecionado('custas')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="🗑️ Remover", command=lambda: self.remover_selecionado('custas')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="📋 Colar do Excel", command=lambda: self.colar_do_excel('custas')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="🧹 Limpar Aba", command=lambda: self.limpar_tree('custas')).pack(side="left", padx=3)

        colunas = ('id', 'descricao', 'data_desembolso', 'valor')
        self.tree_custas = ttk.Treeview(self.tab_custas, columns=colunas, show='headings', height=11)
        titulos = ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico']
        for col, titulo in zip(colunas, titulos):
            self.tree_custas.heading(col, text=titulo)
            self.tree_custas.column(col, width=150, anchor="center")
        self.tree_custas.column('descricao', width=360, anchor="w")
        self.tree_custas.bind("<Double-1>", lambda e: self.editar_selecionado('custas'))

        scroll_y = ttk.Scrollbar(self.tab_custas, orient="vertical", command=self.tree_custas.yview)
        self.tree_custas.configure(yscrollcommand=scroll_y.set)
        scroll_y.pack(side="right", fill="y")
        self.tree_custas.pack(fill="both", expand=True)

    def adicionar_inline_custa(self):
        desc = self.in_custa_desc.get().strip()
        data = self.in_custa_data.get().strip()
        val = self.in_custa_val.get().strip()
        if not data or not val:
            messagebox.showwarning("Campos Obrigatórios", "Preencha a Data e o Valor da custa processual.")
            return
        self.tree_custas.insert('', 'end', values=(
            self.in_custa_id.get().strip(), desc or "Custa Processual", data, val
        ))
        self.in_custa_id.delete(0, 'end')
        self.in_custa_desc.delete(0, 'end')
        self.in_custa_data.delete(0, 'end')
        self.in_custa_val.delete(0, 'end')
        self.in_custa_id.focus_set()
        self._ao_modificar_dados()

    # ================= ABA 4: DEDUÇÕES / BLOQUEIOS (COM INSERÇÃO RÁPIDA) =================
    def montar_aba_deducoes(self):
        f_inline = tk.LabelFrame(
            self.tab_deducoes,
            text=" ⚡ Inserção Rápida de Bloqueio Sisbajud / Depósito Judicial (Digite e pressione [Enter]) ",
            font=("Arial", 9, "bold"), fg="#1D4ED8", padx=8, pady=6
        )
        f_inline.pack(fill="x", pady=(0, 8))

        self.in_ded_id = ttk.Entry(f_inline, width=35)
        self.in_ded_data = ttk.Entry(f_inline, width=18)
        self.in_ded_val = ttk.Entry(f_inline, width=18)

        campos_d = [
            ("ID / Origem (Ex: Evento 128, SISBAJUD1):", self.in_ded_id),
            ("Data do Bloqueio / Depósito:", self.in_ded_data),
            ("Valor Bloqueado / Depositado (R$):", self.in_ded_val)
        ]
        for idx, (rot, w) in enumerate(campos_d):
            tk.Label(f_inline, text=rot, font=("Arial", 8, "bold")).grid(row=0, column=idx, sticky="w", padx=4)
            w.grid(row=1, column=idx, sticky="ew", padx=4, pady=2)
            w.bind("<Return>", lambda e: self.adicionar_inline_deducao())

        btn_inc = tk.Button(
            f_inline, text="➕ Incluir [Enter]", bg="#16A34A", fg="white",
            font=("Arial", 8, "bold"), relief="flat", padx=10, pady=3,
            command=self.adicionar_inline_deducao
        )
        btn_inc.grid(row=1, column=len(campos_d), padx=6, pady=2)

        frame_botoes = ttk.Frame(self.tab_deducoes)
        frame_botoes.pack(fill="x", pady=(0, 6))

        ttk.Button(frame_botoes, text="✏️ Editar Selecionado", command=lambda: self.editar_selecionado('deducoes')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="🗑️ Remover", command=lambda: self.remover_selecionado('deducoes')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="📋 Colar do Excel", command=lambda: self.colar_do_excel('deducoes')).pack(side="left", padx=3)
        ttk.Button(frame_botoes, text="🧹 Limpar Aba", command=lambda: self.limpar_tree('deducoes')).pack(side="left", padx=3)

        colunas = ('id', 'data_bloqueio', 'valor')
        self.tree_deducoes = ttk.Treeview(self.tab_deducoes, columns=colunas, show='headings', height=11)
        titulos = ['ID / Evento / Origem', 'Data do Bloqueio/Depósito', 'Valor (R$)']
        for col, titulo in zip(colunas, titulos):
            self.tree_deducoes.heading(col, text=titulo)
            self.tree_deducoes.column(col, width=200, anchor="center")
        self.tree_deducoes.column('id', width=400, anchor="w")
        self.tree_deducoes.bind("<Double-1>", lambda e: self.editar_selecionado('deducoes'))

        scroll_y = ttk.Scrollbar(self.tab_deducoes, orient="vertical", command=self.tree_deducoes.yview)
        self.tree_deducoes.configure(yscrollcommand=scroll_y.set)
        scroll_y.pack(side="right", fill="y")
        self.tree_deducoes.pack(fill="both", expand=True)

    def adicionar_inline_deducao(self):
        data = self.in_ded_data.get().strip()
        val = self.in_ded_val.get().strip()
        if not data or not val:
            messagebox.showwarning("Campos Obrigatórios", "Preencha a Data e o Valor do bloqueio/depósito.")
            return
        self.tree_deducoes.insert('', 'end', values=(
            self.in_ded_id.get().strip() or "Bloqueio/Depósito", data, val
        ))
        self.in_ded_id.delete(0, 'end')
        self.in_ded_data.delete(0, 'end')
        self.in_ded_val.delete(0, 'end')
        self.in_ded_id.focus_set()
        self._ao_modificar_dados()

    # ================= ABA 5: ACORDO NÃO CUMPRIDO =================
    def montar_aba_acordos(self):
        frame_params = tk.LabelFrame(self.tab_acordos, text=" 📝 Parâmetros da Quebra de Acordo ", font=("Arial", 10, "bold"), padx=10, pady=10)
        frame_params.pack(fill="x", pady=(0, 10))

        campos_acordo = [
            ("Data do Inadimplemento", "Entry", ""),
            ("Regra de Atualização Pós-Quebra", "Combo", self.opcoes_regras),
            ("Valor Original Confessado (Sem Desconto)", "Entry", ""),
            ("Valor do Desconto (A Reverter)", "Entry", ""),
            ("Multa Moratória Contratual (%)", "Entry", "10"),
            ("Honorários de Retomada/Execução (%)", "Entry", "20")
        ]

        row = col = 0
        for nome, tipo, default in campos_acordo:
            var = tk.StringVar()
            var.trace_add("write", lambda *args: self._ao_modificar_dados())
            self.vars_acordo[nome] = var
            lbl = ttk.Label(frame_params, text=f"{nome}:")
            lbl.grid(row=row, column=col, sticky="w", padx=10, pady=5)

            if tipo == "Entry":
                widget = ttk.Entry(frame_params, textvariable=var, width=25)
                if isinstance(default, str) and default:
                    var.set(default)
            elif tipo == "Combo":
                widget = ttk.Combobox(frame_params, textvariable=var, values=default, state="readonly", width=38)
                widget.current(1)

            widget.grid(row=row, column=col+1, sticky="w", padx=10, pady=5)
            col += 2
            if col > 2:
                col = 0
                row += 1

        frame_pagas = tk.LabelFrame(self.tab_acordos, text=" 💸 Parcelas Pagas (A Abater do Saldo) ", font=("Arial", 10, "bold"), padx=10, pady=5)
        frame_pagas.pack(fill="both", expand=True)

        frame_botoes_ac = ttk.Frame(frame_pagas)
        frame_botoes_ac.pack(fill="x", pady=(0, 5))
        ttk.Button(frame_botoes_ac, text="➕ Adicionar Parcela Paga", command=lambda: self.popup_generico('acordo_pagas')).pack(side="left", padx=5)
        ttk.Button(frame_botoes_ac, text="✏️ Editar", command=lambda: self.editar_selecionado('acordo_pagas')).pack(side="left", padx=5)
        ttk.Button(frame_botoes_ac, text="🗑️ Remover", command=lambda: self.remover_selecionado('acordo_pagas')).pack(side="left", padx=5)

        colunas_ac = ('id', 'data_pagamento', 'valor')
        self.tree_acordo_pagas = ttk.Treeview(frame_pagas, columns=colunas_ac, show='headings', height=6)
        titulos_ac = ['Identificação / Parcela', 'Data do Pagamento', 'Valor Pago (R$)']
        for col, titulo in zip(colunas_ac, titulos_ac):
            self.tree_acordo_pagas.heading(col, text=titulo)
            self.tree_acordo_pagas.column(col, width=200, anchor="center")
        self.tree_acordo_pagas.column('id', width=400, anchor="w")

        scroll_y_ac = ttk.Scrollbar(frame_pagas, orient="vertical", command=self.tree_acordo_pagas.yview)
        self.tree_acordo_pagas.configure(yscrollcommand=scroll_y_ac.set)
        scroll_y_ac.pack(side="right", fill="y")
        self.tree_acordo_pagas.pack(fill="both", expand=True)

    # ================= ABA 6: DASHBOARD ESTRATÉGICO =================
    def montar_aba_dashboard(self):
        frame_header = ttk.Frame(self.tab_dashboard)
        frame_header.pack(fill="x", pady=10)
        ttk.Label(frame_header, text="Viabilidade Financeira da Ação", font=("Arial", 14, "bold")).pack(side="left", padx=10)
        ttk.Button(frame_header, text="🔄 Atualizar Gráfico", command=self.atualizar_dashboard).pack(side="right", padx=10)

        self.frame_grafico = ttk.Frame(self.tab_dashboard)
        self.frame_grafico.pack(fill="both", expand=True, pady=20)

        self.canvas_pizza = tk.Canvas(self.frame_grafico, width=300, height=300, bg="#f0f0f0", highlightthickness=0)
        self.canvas_pizza.pack(side="left", padx=30)

        frame_status = ttk.Frame(self.frame_grafico)
        frame_status.pack(side="left", fill="both", expand=True, padx=20)

        self.lbl_status = tk.Label(frame_status, text="Aguardando dados...", font=("Arial", 16, "bold"))
        self.lbl_status.pack(anchor="w", pady=(20, 10))

        self.lbl_detalhes = tk.Label(frame_status, text="", font=("Arial", 12), justify="left")
        self.lbl_detalhes.pack(anchor="w")

    # ================= CONFERÊNCIA RÁPIDA E AUTO-SAVE =================
    def _parse_moeda_ui(self, val):
        return nash.limpar_moeda(val)

    def _ao_modificar_dados(self):
        if hasattr(self, "_timer_resumo") and self._timer_resumo:
            self.root.after_cancel(self._timer_resumo)
        self._timer_resumo = self.root.after(250, self.atualizar_resumo_rapido)

    def atualizar_resumo_rapido(self):
        if not hasattr(self, "tree_danos"):
            return
        qtd_danos = len(self.tree_danos.get_children())
        soma_danos = sum(self._parse_moeda_ui(self.tree_danos.item(i, 'values')[3]) for i in self.tree_danos.get_children() if len(self.tree_danos.item(i, 'values')) > 3)

        qtd_custas = len(self.tree_custas.get_children())
        soma_custas = sum(self._parse_moeda_ui(self.tree_custas.item(i, 'values')[3]) for i in self.tree_custas.get_children() if len(self.tree_custas.item(i, 'values')) > 3)

        qtd_ded = len(self.tree_deducoes.get_children())
        soma_ded = sum(self._parse_moeda_ui(self.tree_deducoes.item(i, 'values')[2]) for i in self.tree_deducoes.get_children() if len(self.tree_deducoes.item(i, 'values')) > 2)

        self.notebook.tab(self.tab_danos, text=f" 2. Valores / Danos ({qtd_danos} | R$ {soma_danos:,.2f}) ")
        self.notebook.tab(self.tab_custas, text=f" 3. Custas ({qtd_custas} | R$ {soma_custas:,.2f}) ")
        self.notebook.tab(self.tab_deducoes, text=f" 4. Bloqueios / Depósitos ({qtd_ded} | R$ {soma_ded:,.2f}) ")

        aviso_datas = ""
        dt_prop = self.vars_param.get("Data Propositura", tk.StringVar()).get().strip()
        dt_tran = self.vars_param.get("Data do Trânsito", tk.StringVar()).get().strip()
        if dt_prop and dt_tran:
            try:
                p_dt = datetime.strptime(dt_prop, "%d/%m/%Y")
                t_dt = datetime.strptime(dt_tran, "%d/%m/%Y")
                if t_dt < p_dt:
                    aviso_datas = "   ⚠️ [Atenção: Data do Trânsito é anterior à Propositura! Se for título extrajudicial, escolha o Modo 2 acima]"
            except Exception:
                pass

        texto = (
            f"Dívida Histórica ({qtd_danos} itens): R$ {soma_danos:,.2f}   -->   "
            f"(+) Custas Históricas ({qtd_custas} guias): R$ {soma_custas:,.2f}   -->   "
            f"(-) Bloqueios/Depósitos ({qtd_ded} itens): -R$ {soma_ded:,.2f}{aviso_datas}"
        )
        self.lbl_resumo_valores.config(text=texto, fg="#B45309" if aviso_datas else "#0F172A")
        self.salvar_rascunho_auto()

    def salvar_rascunho_auto(self):
        try:
            self._gravar_workbook_template(ARQUIVO_RASCUNHO)
            agora = datetime.now().strftime("%H:%M:%S")
            self.lbl_autosave.config(text=f"● Auto-Save Ativo ({agora})")
        except Exception:
            pass

    def _tentar_restaurar_rascunho(self):
        if ARQUIVO_RASCUNHO.exists():
            try:
                self._carregar_workbook_na_tela(ARQUIVO_RASCUNHO, silencioso=True)
            except Exception:
                pass

    def novo_calculo_limpo(self):
        if not messagebox.askyesno("Limpar Tela", "Deseja limpar todos os campos e iniciar um cálculo novo do zero?"):
            return
        self.modo_calculo.set("SENTENCA")
        defaults_limpos = {
            "Processo": "", "Atuação": "AUTOR", "Data do Trânsito": "", "Data da Sentença": "",
            "Termo Inicial Juros": "Desembolso", "Data da Citação": "", "Data do Evento": "",
            "Justiça Gratuita": "Não", "Juros sobre Custas": "Não (Sem Juros)",
            "Pagamento Voluntário 15d": "Não se aplica", "Valor Pagamento Parcial (R$)": "",
            "Fazenda Pública": "Não", "Honorários Sucumbência (%)": "10", "Honorários Fixos (R$)": "",
            "Base Honorários": "CONDENAÇÃO", "Valor Causa Original": "", "Data Propositura": "",
            "Proporção Honorários (%)": "100%", "Proporção Custas (%)": "100%"
        }
        for k, v in defaults_limpos.items():
            if k in self.vars_param:
                self.vars_param[k].set(v)
        for k in self.vars_acordo:
            self.vars_acordo[k].set("10" if "Multa" in k else ("20" if "Honorários" in k else ""))
        for t in ('danos', 'custas', 'deducoes', 'acordo_pagas'):
            self.limpar_tree(t)
        self.aplicar_modo_calculo()
        self.salvar_rascunho_auto()

    def ao_fechar_janela(self):
        self.salvar_rascunho_auto()
        self.root.destroy()

    # ================= ATUALIZADOR AUTOMÁTICO PADRÃO WINDOWS =================
    def _extrair_tupla_versao(self, v_str):
        nums = [int(x) for x in re.findall(r"\d+", str(v_str))]
        return tuple(nums) if nums else (0, 0, 0)

    def verificar_atualizacao_background(self, silencioso=True):
        def _worker():
            try:
                url = f"https://api.github.com/repos/{GITHUB_REPO}/releases/latest"
                req = urllib.request.Request(url, headers={"User-Agent": "NashSystem-Updater"})
                with urllib.request.urlopen(req, timeout=5) as resp:
                    dados = json.loads(resp.read().decode("utf-8"))

                tag_remota = dados.get("tag_name", "")
                versao_local = getattr(nash, "__version__", "4.1.3")
                if self._extrair_tupla_versao(tag_remota) > self._extrair_tupla_versao(versao_local):
                    url_exe = None
                    for asset in dados.get("assets", []):
                        if str(asset.get("name", "")).lower().endswith(".exe"):
                            url_exe = asset.get("browser_download_url")
                            break
                    url_html = dados.get("html_url", f"https://github.com/{GITHUB_REPO}/releases")
                    self.root.after(0, lambda: self._oferecer_atualizacao(tag_remota, versao_local, url_exe, url_html))
                elif not silencioso:
                    self.root.after(0, lambda: messagebox.showinfo(
                        "Sistema Atualizado",
                        f"Você já está utilizando a versão mais recente do Nash System (v{versao_local})."
                    ))
            except Exception as e:
                if not silencioso:
                    self.root.after(0, lambda: messagebox.showwarning(
                        "Verificação de Atualização",
                        f"Não foi possível consultar atualizações no momento:\n{e}"
                    ))

        threading.Thread(target=_worker, daemon=True).start()

    def _oferecer_atualizacao(self, tag_remota, versao_local, url_exe, url_html):
        resp = messagebox.askyesno(
            "Atualização Automática — Nash System",
            f"Uma nova versão do Nash System ({tag_remota}) está disponível!\n"
            f"(Versão instalada: v{versao_local})\n\n"
            "O sistema irá baixar e instalar a atualização automaticamente em segundo plano "
            "(não exige desinstalar e mantém o seu rascunho atual salvo).\n\n"
            "Deseja atualizar agora?"
        )
        if not resp:
            return
        self.salvar_rascunho_auto()
        if url_exe and sys.platform.startswith("win"):
            self._baixar_e_executar_instalador(url_exe)
        else:
            import webbrowser
            webbrowser.open(url_html)

    def _baixar_e_executar_instalador(self, url_exe):
        progresso_win = tk.Toplevel(self.root)
        progresso_win.title("Atualizando Nash System...")
        progresso_win.geometry("420x130")
        progresso_win.transient(self.root)
        progresso_win.grab_set()
        tk.Label(progresso_win, text="Baixando atualização automática... Aguarde alguns segundos.", font=("Arial", 10, "bold")).pack(pady=15)
        barra = ttk.Progressbar(progresso_win, mode="indeterminate", length=360)
        barra.pack(pady=10)
        barra.start(10)

        def _download():
            try:
                destino = Path(tempfile.gettempdir()) / "Instalador_Nash_Update.exe"
                req = urllib.request.Request(url_exe, headers={"User-Agent": "NashSystem-Updater"})
                with urllib.request.urlopen(req, timeout=60) as r, open(destino, "wb") as f:
                    f.write(r.read())
                subprocess.Popen([str(destino), "/SILENT", "/CLOSEAPPLICATIONS", "/RESTARTAPPLICATIONS"])
                self.root.after(0, self.root.destroy)
            except Exception as e:
                self.root.after(0, lambda: [
                    progresso_win.destroy(),
                    messagebox.showerror("Erro ao Atualizar", f"Falha ao baixar instalador:\n{e}")
                ])

        threading.Thread(target=_download, daemon=True).start()

    # ================= DASHBOARD E TABELAS =================
    def atualizar_dashboard(self):
        total_pedido = 0.0
        for item in self.tree_danos.get_children():
            valores = self.tree_danos.item(item, 'values')
            str_pedido = valores[6] if len(valores) > 6 and str(valores[6]).strip() else (valores[3] if len(valores) > 3 else "0")
            total_pedido += self._parse_moeda_ui(str_pedido)

        total_custas = 0.0
        for item in self.tree_custas.get_children():
            valores = self.tree_custas.item(item, 'values')
            str_custas = valores[3] if len(valores) > 3 and str(valores[3]).strip() else "0"
            total_custas += self._parse_moeda_ui(str_custas)

        if total_pedido <= 0:
            self.lbl_status.config(text="Sem dados de pedido inicial para análise.", fg="gray")
            self.canvas_pizza.delete("all")
            self.lbl_detalhes.config(text="")
            return

        proporcao = total_custas / total_pedido
        atuacao = self.vars_param.get("Atuação", tk.StringVar()).get().strip().upper()
        jg_ativa = self.vars_param.get("Justiça Gratuita", tk.StringVar()).get().strip() == "Sim"
        parte_adversa = "Autor" if atuacao == "RÉU" else "Réu"

        self.canvas_pizza.delete("all")
        angulo_custas = min(360, (proporcao * 360))
        angulo_restante = max(0, 360 - angulo_custas)

        if proporcao >= 0.70:
            cor_custas = "#FF4C4C"
            texto_status = "🔴 INVIABILIDADE ECONÔMICA"
        elif proporcao >= 0.50:
            cor_custas = "#FFD700"
            texto_status = "🟡 ATENÇÃO: RISCO FINANCEIRO"
        else:
            cor_custas = "#4CAF50"
            texto_status = "🟢 OPERAÇÃO SAUDÁVEL"

        self.canvas_pizza.create_arc(10, 10, 290, 290, start=90, extent=-angulo_custas, fill=cor_custas, outline="white", width=2)
        self.canvas_pizza.create_arc(10, 10, 290, 290, start=90-angulo_custas, extent=-angulo_restante, fill="#2196F3", outline="white", width=2)

        self.lbl_status.config(text=texto_status, fg=cor_custas)
        texto_detalhes = (
            f"Valor do Pedido Original: R$ {total_pedido:,.2f}\n"
            f"Custas Adiantadas: R$ {total_custas:,.2f}\n"
            f"Comprometimento: {proporcao*100:.2f}%\n\n"
        )
        if jg_ativa:
            texto_detalhes += f"⚠ Oponente ({parte_adversa}) possui Justiça Gratuita.\nMesmo com êxito, o desembolso de custas é um Sunk Cost (irrecuperável)."
        else:
            texto_detalhes += f"Oponente ({parte_adversa}) NÃO possui Justiça Gratuita.\nHá expectativa legal de recuperar as custas no fim do processo."

        self.lbl_detalhes.config(text=texto_detalhes)

    def colar_do_excel(self, tipo_tree):
        try:
            texto = self.root.clipboard_get()
        except Exception:
            messagebox.showwarning("Área de Transferência Vazia", "Copie as células no Excel (Ctrl+C) antes de clicar em Colar.")
            return

        linhas = [l for l in texto.splitlines() if l.strip()]
        if not linhas:
            return

        tree_alvo = getattr(self, f"tree_{tipo_tree}")
        qtd_cols = len(tree_alvo["columns"])
        incluidos = 0
        for l in linhas:
            partes = [p.strip() for p in l.split("\t")]
            if len(partes) < 2:
                continue
            valores = (partes + [""] * qtd_cols)[:qtd_cols]
            if tipo_tree == "danos" and not valores[4]:
                valores[4] = "R1"
            tree_alvo.insert("", "end", values=valores)
            incluidos += 1

        self._ao_modificar_dados()
        if incluidos:
            messagebox.showinfo("Colagem Concluída", f"{incluidos} linha(s) colada(s) com sucesso!")

    # ================= LÓGICA DE POPUPS UNIFICADA =================
    def popup_generico(self, tipo_tree, item_selecionado=None):
        popup = tk.Toplevel(self.root)
        popup.transient(self.root)
        popup.grab_set()

        if tipo_tree == 'danos':
            popup.title("Editar Dano" if item_selecionado else "Adicionar Dano")
            campos = ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico', 'Regra', 'Data Juros', 'Valor Ped. Inicial', 'Data do Pedido']
            tree_alvo = self.tree_danos
        elif tipo_tree == 'custas':
            popup.title("Editar Custa" if item_selecionado else "Adicionar Custa")
            campos = ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico']
            tree_alvo = self.tree_custas
        elif tipo_tree == 'deducoes':
            popup.title("Editar Dedução" if item_selecionado else "Adicionar Depósito/Bloqueio")
            campos = ['ID / Folha', 'Data bloqueio/deposito', 'Valor']
            tree_alvo = self.tree_deducoes
        elif tipo_tree == 'acordo_pagas':
            popup.title("Editar Parcela Paga" if item_selecionado else "Adicionar Parcela Paga")
            campos = ['Identificação / Parcela', 'Data do Pagamento', 'Valor Pago (R$)']
            tree_alvo = self.tree_acordo_pagas

        entradas = {}
        valores_atuais = tree_alvo.item(item_selecionado, 'values') if item_selecionado else [''] * len(campos)

        for i, campo in enumerate(campos):
            ttk.Label(popup, text=f"{campo}:").grid(row=i, column=0, padx=10, pady=5, sticky="w")

            if campo == 'Regra':
                ent = ttk.Combobox(popup, values=self.opcoes_regras, state="readonly", width=37)
                val_atual = str(valores_atuais[i]) if i < len(valores_atuais) and valores_atuais[i] else ""
                encontrou = False
                for opt in self.opcoes_regras:
                    if opt.startswith(val_atual) and val_atual:
                        ent.set(opt)
                        encontrou = True
                        break
                if not encontrou:
                    ent.current(0)
            else:
                ent = ttk.Entry(popup, width=40)
                if i < len(valores_atuais) and valores_atuais[i]:
                    ent.insert(0, valores_atuais[i])

            ent.grid(row=i, column=1, padx=10, pady=5)
            entradas[campo] = ent

        def salvar():
            novos_valores = []
            for c in campos:
                val = entradas[c].get()
                if c == 'Regra' and " - " in val:
                    val = val.split(" - ")[0]
                novos_valores.append(val)

            if item_selecionado:
                tree_alvo.item(item_selecionado, values=novos_valores)
            else:
                tree_alvo.insert('', 'end', values=novos_valores)
            self._ao_modificar_dados()
            popup.destroy()

        ttk.Button(popup, text="Salvar Alterações" if item_selecionado else "Salvar Linha", command=salvar).grid(row=len(campos), column=0, columnspan=2, pady=15)

    def editar_selecionado(self, tipo_tree):
        tree = getattr(self, f"tree_{tipo_tree}")
        selecionados = tree.selection()
        if not selecionados:
            messagebox.showwarning("Aviso", "Selecione uma linha para editar.")
            return
        self.popup_generico(tipo_tree, item_selecionado=selecionados[0])

    def remover_selecionado(self, tipo_tree):
        tree = getattr(self, f"tree_{tipo_tree}")
        for item in tree.selection():
            tree.delete(item)
        self._ao_modificar_dados()

    def limpar_tree(self, tipo_tree):
        tree = getattr(self, f"tree_{tipo_tree}")
        for item in tree.get_children():
            tree.delete(item)
        self._ao_modificar_dados()

    # ================= EXPORTAÇÃO E IMPORTAÇÃO DE TEMPLATES =================
    def _carregar_workbook_na_tela(self, caminho, silencioso=False):
        wb = load_workbook(caminho, data_only=True)

        if 'Parametros' in wb.sheetnames:
            ws = wb['Parametros']
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    chave = str(row[0]).strip()
                    valor = row[1] if len(row) > 1 and row[1] is not None else ""
                    if isinstance(valor, datetime):
                        valor = valor.strftime("%d/%m/%Y")
                    if chave in self.vars_param:
                        self.vars_param[chave].set(str(valor))

        def carregar_tree(sheet_name, tree_widget, cols_esperadas):
            if sheet_name in wb.sheetnames:
                for item in tree_widget.get_children():
                    tree_widget.delete(item)
                ws = wb[sheet_name]
                headers = [cell.value for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not any(row):
                        continue
                    row_data = dict(zip(headers, row))
                    valores = []
                    for col in cols_esperadas:
                        val = row_data.get(col, "")
                        val = val if val is not None else ""
                        if isinstance(val, datetime):
                            val = val.strftime("%d/%m/%Y")
                        valores.append(str(val))
                    if any(v.strip() for v in valores):
                        tree_widget.insert('', 'end', values=valores)

        carregar_tree('Danos', self.tree_danos, ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico', 'Regra', 'Data Juros', 'Valor Pedido Inicial', 'Data do Pedido'])
        carregar_tree('Custas', self.tree_custas, ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico'])
        carregar_tree('Deducoes', self.tree_deducoes, ['ID / Folha', 'Data bloqueio/deposito', 'Valor'])

        if 'Acordo_Params' in wb.sheetnames:
            ws = wb['Acordo_Params']
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    chave = str(row[0]).strip()
                    valor = row[1] if len(row) > 1 and row[1] is not None else ""
                    if isinstance(valor, datetime):
                        valor = valor.strftime("%d/%m/%Y")
                    if chave in self.vars_acordo:
                        self.vars_acordo[chave].set(str(valor))

        carregar_tree('Acordo_Pagas', self.tree_acordo_pagas, ['Identificação / Parcela', 'Data do Pagamento', 'Valor Pago (R$)'])

        self.root.after(100, self.atualizar_dashboard)
        self.root.after(100, self.atualizar_resumo_rapido)
        if not silencioso:
            messagebox.showinfo("Sucesso", "Planilha carregada com sucesso!")

    def importar_planilha(self):
        caminho = filedialog.askopenfilename(title="Selecione o Template Excel", filetypes=[("Planilha Excel", "*.xlsx")])
        if not caminho:
            return
        try:
            self._carregar_workbook_na_tela(caminho, silencioso=False)
        except Exception as e:
            messagebox.showerror("Erro ao Importar", f"Ocorreu um erro ao tentar ler a planilha:\n{str(e)}")

    def calcular_direto_planilha(self):
        """Permite escolher um arquivo template.xlsx editado no Excel e gerar o Laudo + PDF direto."""
        caminho_in = filedialog.askopenfilename(title="Selecione o Template preenchido no Excel", filetypes=[("Planilha Excel", "*.xlsx")])
        if not caminho_in:
            return
        try:
            self._carregar_workbook_na_tela(caminho_in, silencioso=True)
            p_in = Path(caminho_in)
            nome_sugerido = p_in.name.replace("template ", "laudo ").replace("Template ", "laudo ")
            if nome_sugerido == p_in.name:
                nome_sugerido = f"laudo_{p_in.name}"
            caminho_out = filedialog.asksaveasfilename(
                title="Onde salvar o Laudo de Liquidação?",
                initialdir=str(p_in.parent),
                initialfile=nome_sugerido,
                defaultextension=".xlsx",
                filetypes=[("Planilha Excel", "*.xlsx")]
            )
            if not caminho_out:
                return
            self._disparar_thread_motor(Path(caminho_in), Path(caminho_out))
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível processar a planilha selecionada:\n{e}")

    def _gravar_workbook_template(self, caminho_arquivo):
        wb = Workbook()
        f_bold = Font(bold=True)
        borda = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        fill_cinza = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

        def salvar_params(aba, dicionario):
            ws = wb.create_sheet(aba) if aba != 'Parametros' else wb.active
            ws.title = aba
            for r_idx, (chave, var_tk) in enumerate(dicionario.items(), 1):
                val = var_tk.get()
                if 'Regra' in chave and ' - ' in val:
                    val = val.split(' - ')[0]
                c1 = ws.cell(row=r_idx, column=1, value=chave)
                c1.font = f_bold
                c1.fill = fill_cinza
                c1.border = borda
                ws.cell(row=r_idx, column=2, value=val).border = borda
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 30

        def salvar_tree(aba, tree_widget, headers):
            ws = wb.create_sheet(aba)
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.font = f_bold
                cell.fill = fill_cinza
                cell.border = borda
                ws.column_dimensions[chr(64+col_idx)].width = 20
            for row_idx, item in enumerate(tree_widget.get_children(), 2):
                for col_idx, val in enumerate(tree_widget.item(item, 'values'), 1):
                    ws.cell(row=row_idx, column=col_idx, value=val).border = borda

        salvar_params('Parametros', self.vars_param)
        salvar_tree('Danos', self.tree_danos, ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico', 'Regra', 'Data Juros', 'Valor Pedido Inicial', 'Data do Pedido'])
        salvar_tree('Custas', self.tree_custas, ['ID / Folha', 'Descrição', 'Data Desembolso', 'Valor Histórico'])
        salvar_tree('Deducoes', self.tree_deducoes, ['ID / Folha', 'Data bloqueio/deposito', 'Valor'])
        salvar_params('Acordo_Params', self.vars_acordo)
        salvar_tree('Acordo_Pagas', self.tree_acordo_pagas, ['Identificação / Parcela', 'Data do Pagamento', 'Valor Pago (R$)'])

        wb.save(caminho_arquivo)

    def salvar_apenas_template(self):
        processo_bruto = self.vars_param["Processo"].get().strip() or "sem_numero"
        processo_seguro = "".join(c for c in processo_bruto if c not in r'\/:*?"<>|')
        data_atual = datetime.now().strftime("%d.%m.%Y")

        caminho = filedialog.asksaveasfilename(
            title="Salvar Template Excel",
            defaultextension=".xlsx",
            initialfile=f"template {processo_seguro} {data_atual}.xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")]
        )
        if not caminho:
            return
        try:
            self._gravar_workbook_template(Path(caminho))
            messagebox.showinfo("Template Salvo", f"Template guardado com sucesso em:\n{caminho}")
        except Exception as e:
            messagebox.showerror("Erro ao Salvar Template", str(e))

    def salvar_e_gerar(self):
        try:
            processo_bruto = self.vars_param["Processo"].get().strip() or "sem_numero"
            processo_seguro = "".join(c for c in processo_bruto if c not in r'\/:*?"<>|')
            data_atual = datetime.now().strftime("%d.%m.%Y")

            caminho_destino = filedialog.asksaveasfilename(
                title="Salvar Laudo de Liquidação",
                defaultextension=".xlsx",
                initialfile=f"laudo {processo_seguro} {data_atual}.xlsx",
                filetypes=[("Planilha Excel", "*.xlsx"), ("Todos os arquivos", "*.*")]
            )
            if not caminho_destino:
                return

            path_destino = Path(caminho_destino)
            pasta_destino = path_destino.parent
            nome_base = path_destino.stem
            nome_template = pasta_destino / f"template {processo_seguro} {data_atual}.xlsx"
            nome_laudo_xls = pasta_destino / f"{nome_base}.xlsx"

            self._gravar_workbook_template(nome_template)
            self._disparar_thread_motor(nome_template, nome_laudo_xls)

        except Exception as e:
            messagebox.showerror("Erro Crítico", f"Falha ao preparar o Excel:\n{str(e)}")

    def _disparar_thread_motor(self, nome_template, nome_laudo_xls):
        self.btn_processar.config(state="disabled")
        pasta_destino = Path(nome_laudo_xls).parent

        def thread_calculo():
            try:
                resultado = nash.executar_nash(str(nome_template), str(nome_laudo_xls))
                def sucesso():
                    msg_base = (
                        f"Processamento finalizado com sucesso!\n\n"
                        f"• Template salvo: {Path(nome_template).name}\n"
                        f"• Laudo salvo: {Path(nome_laudo_xls).name}\n"
                        f"• Pasta: {pasta_destino}"
                    )
                    if isinstance(resultado, dict) and resultado.get("nivel") in ["ALERTA", "CRITICO"]:
                        msg_alerta = resultado.get("mensagem", "")
                        messagebox.showwarning("Aviso de Viabilidade Financeira", f"{msg_base}\n\n⚠️ {msg_alerta}")
                    else:
                        messagebox.showinfo("Cálculo Concluído", msg_base)
                    self.btn_processar.config(state="normal")
                self.root.after(0, sucesso)
            except Exception as e:
                erro_msg = str(e)
                def falha():
                    messagebox.showerror("Erro no Motor Matemático", f"O cálculo falhou:\n\n{erro_msg}")
                    self.btn_processar.config(state="normal")
                self.root.after(0, falha)

        threading.Thread(target=thread_calculo, daemon=True).start()


if __name__ == "__main__":
    root = tk.Tk()
    style = ttk.Style()
    if 'clam' in style.theme_names():
        style.theme_use('clam')
    app = NashDataEntryGUI(root)
    root.mainloop()